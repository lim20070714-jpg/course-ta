from pathlib import Path

from pptx import Presentation
from pypdf import PdfReader
from sqlmodel import Session

from app import db
from app.config import settings
from app.core import llm
from app.core.store import to_blob
from app.models import Chunk, Material

_EMBED_BATCH = 16


def extract_pages(path: str | Path) -> list[tuple[int, str]]:
    """提取文档每一页/每张幻灯片的文本，返回 [(页码, 文本), ...]。"""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(p)
    if suffix == ".pptx":
        return _extract_pptx(p)
    if suffix in (".txt", ".md"):
        return [(1, p.read_text(encoding="utf-8", errors="ignore"))]
    if suffix == ".ppt":
        raise ValueError("旧版 .ppt 不受支持，请另存为 .pptx 后重新上传")
    raise ValueError(f"不支持的文件类型: {suffix}")


def _extract_pdf(p: Path) -> list[tuple[int, str]]:
    reader = PdfReader(str(p))
    return [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]


def _extract_pptx(p: Path) -> list[tuple[int, str]]:
    prs = Presentation(str(p))
    pages: list[tuple[int, str]] = []
    for i, slide in enumerate(prs.slides):
        parts = [
            shape.text_frame.text.strip()
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(notes)
        pages.append((i + 1, "\n".join(parts)))
    return pages


def chunk_pages(
    pages: list[tuple[int, str]], chunk_size: int | None = None, overlap: int | None = None
) -> list[dict]:
    """把每页文本切成带重叠的片段，保留页码信息。"""
    chunk_size = chunk_size or settings.chunk_size
    overlap = overlap or settings.chunk_overlap
    chunks: list[dict] = []
    seq = 0
    for page_no, text in pages:
        text = " ".join(text.split())
        start = 0
        while start < len(text):
            piece = text[start : start + chunk_size].strip()
            if piece:
                chunks.append({"page": page_no, "seq": seq, "text": piece})
                seq += 1
            if start + chunk_size >= len(text):
                break
            start += chunk_size - overlap
    return chunks


def load_and_chunk(path: str | Path) -> list[dict]:
    return chunk_pages(extract_pages(path))


async def ingest_material(material_id: int) -> None:
    """后台任务：解析材料 -> 分块 -> 向量化 -> 入库。"""
    with Session(db.engine) as session:
        material = session.get(Material, material_id)
        if material is None:
            return
        material.status = "ingesting"
        session.add(material)
        session.commit()
        try:
            chunks = load_and_chunk(material.stored_path)
            if not chunks:
                raise ValueError("未从文件中提取到文本内容（扫描版 PDF 需要先做 OCR）")
            vectors: list[list[float]] = []
            for i in range(0, len(chunks), _EMBED_BATCH):
                batch = [c["text"] for c in chunks[i : i + _EMBED_BATCH]]
                vectors.extend(await llm.embed(batch))
            for c, vec in zip(chunks, vectors, strict=True):
                session.add(
                    Chunk(
                        course_id=material.course_id,
                        material_id=material.id,
                        page=c["page"],
                        seq=c["seq"],
                        text=c["text"],
                        embedding=to_blob(vec),
                    )
                )
            material.chunk_count = len(chunks)
            material.status = "ready"
            material.error = ""
        except Exception as exc:  # noqa: BLE001  把失败原因暴露给上传者
            material.status = "failed"
            material.error = str(exc)[:500]
        session.add(material)
        session.commit()

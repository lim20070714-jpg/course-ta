"""Word 文档与图片材料：解析、内嵌图片导出、页图服务与对外字段。"""

import asyncio
import json
import sqlite3
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from sqlmodel import Session, SQLModel, create_engine
from sqlmodel.pool import StaticPool

import app.db as db_module
from app.config import settings
from app.core import media
from app.core.ingest import extract_docx_pages, extract_pages, ingest_pages
from app.core.kinds import detect_kind
from app.main import app
from app.models import Course, Material

# ---------------------------------------------------------------- 类型识别


def test_detect_kind_recognises_images_by_suffix() -> None:
    for name in ("电路图.png", "photo.JPG", "scan.jpeg", "波形.webp", "笔记.tiff"):
        assert detect_kind(name) == "image", name


def test_detect_kind_recognises_screenshot_names() -> None:
    assert detect_kind("期中截图-第3章.pdf") == "image"
    assert detect_kind("screenshot-2026.png") == "image"


def test_detect_kind_existing_rules_unchanged() -> None:
    # 既有识别规则不受影响
    assert detect_kind("016秋A末4.pdf") == "exam"
    assert detect_kind("A1-01电路分析.pdf") == "slides"
    assert detect_kind("电电1A小班辅导-真题讲解-含批注.pdf") == "exam"


# ---------------------------------------------------------------- 图片工具


def _png(path: Path, size: tuple[int, int] = (40, 30), color: tuple[int, int, int] = (200, 40, 40)) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", size, color).save(path)
    return path


def test_image_map_parses_relative_and_absolute(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    absolute = tmp_path / "abs.png"
    raw = json.dumps(
        [
            {"page": 3, "path": "media/7/docx/image-001.png"},
            {"page": 5, "path": str(absolute)},
            {"page": "bad", "path": "x.png"},
            {"nope": 1},
        ]
    )
    mapping = media.image_map(raw)
    assert set(mapping) == {3, 5}
    assert mapping[3] == str(tmp_path / "media/7/docx/image-001.png")
    assert mapping[5] == str(absolute)
    assert media.image_page_numbers(raw) == [3, 5]


def test_image_map_handles_garbage() -> None:
    for raw in ("", "[]", "not json", "{}", '{"a":1}'):
        assert media.image_map(raw) == {}


def test_write_images_sniffs_real_format(tmp_path: Path) -> None:
    png = _png(tmp_path / "a" / "x.png").read_bytes()
    written = media.write_images([png], tmp_path / "out")
    assert len(written) == 1
    assert written[0].suffix == ".png"
    # 内容是 JPEG 时按文件头纠正扩展名
    jpg_path = tmp_path / "b" / "y.jpg"
    jpg_path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (20, 20)).save(jpg_path, format="JPEG")
    written_jpg = media.write_images([jpg_path.read_bytes()], tmp_path / "out2")
    assert written_jpg[0].suffix == ".jpg"


def test_media_type_for() -> None:
    assert media.media_type_for("a.png") == "image/png"
    assert media.media_type_for("a.JPG") == "image/jpeg"
    assert media.media_type_for("a.webp") == "image/webp"
    assert media.media_type_for("a.unknown") == "image/png"


def test_render_page_png_and_embedded_pdf_page(tmp_path: Path) -> None:
    # 用一个真实 PDF 验证页图渲染仍可用（由 PDF 材料自己生成）
    pdf = tmp_path / "t.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "hello")
    doc.save(str(pdf))
    doc.close()
    png = media.render_page_png(pdf, 1)
    assert png[:8] == b"\x89PNG\r\n\x1a\n"
    with pytest.raises(ValueError):
        media.render_page_png(pdf, 99)


# ---------------------------------------------------------------- docx 解析


def _make_docx(path: Path, paragraphs: list[str], image: Path | None = None) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if image is not None:
        doc.add_picture(str(image))
    doc.save(str(path))
    return path


def test_extract_docx_pages_returns_text(tmp_path: Path) -> None:
    docx = _make_docx(tmp_path / "笔记.docx", ["第一章 电路分析", "戴维南定理：$R_{th}$"])
    pages = extract_docx_pages(docx)
    assert pages == [(1, "第一章 电路分析\n戴维南定理：$R_{th}$")]


def test_extract_pages_dispatch_docx_and_image(tmp_path: Path) -> None:
    docx = _make_docx(tmp_path / "a.docx", ["内容"])
    assert extract_pages(docx) == [(1, "内容")]
    png = _png(tmp_path / "b.png")
    assert extract_pages(png) == [(1, "")]  # 图片本身没有文字，正文来自视觉描述


def test_extract_pages_rejects_legacy_formats(tmp_path: Path) -> None:
    for name in ("old.doc", "old.ppt"):
        p = tmp_path / name
        p.write_bytes(b"x")
        with pytest.raises(ValueError, match="另存为"):
            extract_pages(p)


def test_extract_pages_rejects_unknown(tmp_path: Path) -> None:
    p = tmp_path / "a.zip"
    p.write_bytes(b"x")
    with pytest.raises(ValueError, match="不支持"):
        extract_pages(p)


def test_extract_docx_body_returns_embedded_image(tmp_path: Path) -> None:
    image = _png(tmp_path / "src.png", (24, 24))
    docx = _make_docx(tmp_path / "含图.docx", ["正文"], image=image)
    text, blobs = media.extract_docx_body(docx)
    assert "正文" in text
    assert len(blobs) == 1
    assert blobs[0].startswith(b"\x89PNG")


def test_extract_docx_body_reads_tables(tmp_path: Path) -> None:
    path = tmp_path / "表.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "元件"
    table.rows[0].cells[1].text = "阻值"
    doc.save(str(path))
    text, _ = media.extract_docx_body(path)
    assert "元件 | 阻值" in text


# ---------------------------------------------------------------- ingest_pages


def test_ingest_pages_docx_keeps_body_and_exports_images(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "data_dir", tmp_path / "data")
    monkeypatch.setattr(settings, "vision_model", "")  # 不调用视觉模型
    image = _png(tmp_path / "pic.png")
    docx = _make_docx(tmp_path / "讲义.docx", ["第一章", "第二章"], image=image)

    pages, images = asyncio.run(ingest_pages(docx, 42))

    assert pages == [(1, "第一章\n第二章")]
    assert len(images) == 1
    page_no, rel = images[0]
    assert page_no == 2  # 内嵌图片接在正文之后
    stored = tmp_path / "data" / rel
    assert stored.exists() and stored.read_bytes().startswith(b"\x89PNG")


def test_ingest_pages_image_material_is_page_one(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "vision_model", "")
    png = _png(tmp_path / "电路图.png")
    pages, images = asyncio.run(ingest_pages(png, 7))
    assert pages == []
    assert images == [(1, str(png))]


def test_ingest_pages_vision_disabled_skips_description(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "vision_model", "")
    docx = _make_docx(tmp_path / "a.docx", ["正文"], image=_png(tmp_path / "i.png"))
    pages, _images = asyncio.run(ingest_pages(docx, 1))
    assert all(page == 1 for page, _ in pages)  # 没有描述就不产生图片页文字


def test_ingest_pages_pptx_exports_images(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "data_dir", tmp_path / "data")
    monkeypatch.setattr(settings, "vision_model", "")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "第一讲"
    image = _png(tmp_path / "s.png", (32, 32))
    slide.shapes.add_picture(str(image), Inches(1), Inches(1), Inches(2), Inches(2))
    pptx = tmp_path / "课件.pptx"
    prs.save(str(pptx))

    pages, images = asyncio.run(ingest_pages(pptx, 9))
    assert pages and pages[0][1].strip() == "第一讲"
    assert len(images) == 1
    assert (tmp_path / "data" / images[0][1]).exists()


# ---------------------------------------------------------------- 数据库迁移与 API


def test_migration_adds_pages_column(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:

    from sqlmodel import create_engine

    legacy = tmp_path / "legacy.db"
    con = sqlite3.connect(legacy)
    con.execute(
        "CREATE TABLE material (id INTEGER PRIMARY KEY, course_id INTEGER, filename VARCHAR NOT NULL,"
        " stored_path VARCHAR NOT NULL, kind VARCHAR NOT NULL DEFAULT 'other',"
        " status VARCHAR NOT NULL DEFAULT 'pending', error VARCHAR NOT NULL DEFAULT '',"
        " chunk_count INTEGER NOT NULL DEFAULT 0, created_at DATETIME NOT NULL)"
    )
    con.execute(
        "INSERT INTO material (course_id, filename, stored_path, created_at)"
        " VALUES (1, '老材料.pdf', 'x', '2026-01-01')"
    )
    con.commit()
    con.close()

    db_module.engine = create_engine(f"sqlite:///{legacy.as_posix()}")
    db_module.init_db()
    con = sqlite3.connect(legacy)
    cols = [r[1] for r in con.execute("PRAGMA table_info(material)")]
    assert "pages" in cols
    assert con.execute("SELECT pages FROM material").fetchone()[0] == "[]"
    assert con.execute("SELECT filename FROM material").fetchone()[0] == "老材料.pdf"
    db_module.init_db()  # 幂等


@pytest.fixture(name="client")
def client_fixture(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> TestClient:
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    SQLModel.metadata.create_all(engine)
    monkeypatch.setattr(db_module, "engine", engine)
    monkeypatch.setattr(settings, "data_dir", tmp_path / "data")
    with Session(engine) as session:
        session.add(Course(id=1, name="电路"))
        session.commit()
    return TestClient(app)


def test_upload_rejects_unsupported_with_format_hint(client: TestClient) -> None:
    resp = client.post(
        "/api/courses/1/materials",
        files={"file": ("a.zip", b"x", "application/zip")},
        data={"kind": "auto"},
    )
    assert resp.status_code == 400
    assert "DOCX" in resp.json()["detail"] and "图片" in resp.json()["detail"]


def test_upload_accepts_image(client: TestClient, tmp_path: Path) -> None:
    png = _png(tmp_path / "电路图.png").read_bytes()
    resp = client.post(
        "/api/courses/1/materials",
        files={"file": ("电路图.png", png, "image/png")},
        data={"kind": "auto"},
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["kind"] == "image"
    assert body["filename"] == "电路图.png"


def test_upload_accepts_docx(client: TestClient, tmp_path: Path) -> None:
    docx = _make_docx(tmp_path / "讲义.docx", ["第一章"])
    resp = client.post(
        "/api/courses/1/materials",
        files={
            "file": (
                "讲义.docx",
                docx.read_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={"kind": "auto"},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["filename"] == "讲义.docx"


def test_page_image_serves_stored_image(client: TestClient, tmp_path: Path) -> None:
    """图片材料的原页图直接返回该文件（不是 PDF 渲染）。"""
    target = tmp_path / "data" / "materials" / "1" / "电路图.png"
    _png(target, (36, 24))
    with Session(db_module.engine) as session:
        material = Material(
            course_id=1,
            filename="电路图.png",
            stored_path=str(target),
            kind="image",
            status="ready",
            pages=json.dumps([{"page": 1, "path": "materials/1/电路图.png"}]),
        )
        session.add(material)
        session.commit()
        session.refresh(material)
        mid = material.id

    resp = client.get(f"/api/materials/{mid}/page-image?page=1")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("image/png")
    assert resp.content[:8] == b"\x89PNG\r\n\x1a\n"


def test_page_image_reports_missing_file(client: TestClient) -> None:
    with Session(db_module.engine) as session:
        material = Material(
            course_id=1,
            filename="丢了.png",
            stored_path="nowhere.png",
            kind="image",
            status="ready",
            pages=json.dumps([{"page": 1, "path": "materials/1/gone.png"}]),
        )
        session.add(material)
        session.commit()
        session.refresh(material)
        mid = material.id
    resp = client.get(f"/api/materials/{mid}/page-image?page=1")
    assert resp.status_code == 404
    assert "丢失" in resp.json()["detail"]


def test_material_list_exposes_pages_and_image_pages(client: TestClient) -> None:
    with Session(db_module.engine) as session:
        session.add(
            Material(
                course_id=1,
                filename="含图.docx",
                stored_path="x.docx",
                kind="handout",
                status="ready",
                chunk_count=3,
                pages=json.dumps([{"page": 4, "path": "media/1/docx/image-001.png"}]),
            )
        )
        session.commit()
    row = client.get("/api/courses/1/materials").json()[0]
    assert row["image_pages"] == [4]
    assert row["page_count"] == 4


def test_non_pdf_without_image_page_is_rejected(client: TestClient) -> None:
    with Session(db_module.engine) as session:
        material = Material(
            course_id=1, filename="a.docx", stored_path="a.docx", kind="handout", status="ready"
        )
        session.add(material)
        session.commit()
        session.refresh(material)
        mid = material.id
    resp = client.get(f"/api/materials/{mid}/page-image?page=2")
    assert resp.status_code == 400
    assert "不是 PDF" in resp.json()["detail"]
